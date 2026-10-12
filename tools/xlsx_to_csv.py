#!/usr/bin/env python3
"""把 .xlsx 的第一个工作表导出成 CSV，供 build_data.js 读取。

用法： python xlsx_to_csv.py <输入.xlsx> <输出.csv>
"""
import csv
import sys

try:
    import openpyxl
except ImportError:
    sys.stderr.write("缺少 openpyxl，无法读取 .xlsx。请改用 CSV，或用内置 Python 运行。\n")
    sys.exit(2)


def main():
    if len(sys.argv) != 3:
        sys.stderr.write(__doc__)
        sys.exit(1)

    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    with open(dst, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            writer.writerow(["" if c is None else c for c in row])

    wb.close()


if __name__ == "__main__":
    main()

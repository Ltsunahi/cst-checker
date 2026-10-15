#!/usr/bin/env python3
"""把「报名表 + 评分」的合并表转成站点需要的 学号,姓名,结果 三列 CSV。

用法：
  python xlsx_scores_to_csv.py <输入.xlsx> <输出.csv> [--fail-max N] [--ref 参考表.xlsx]

  --fail-max N  评分 <= N 判为「未通过」，默认 2
  --ref         缺学号时，从这张表的「姓名 / 学号」两列里回填（可选）

列名按关键字模糊匹配，所以「人名/姓名」「评分（5分制）/分数/得分」这类都能认出来。
"""
import argparse
import csv
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.stderr.write("缺少 openpyxl，请用内置 Python 运行。\n")
    sys.exit(2)

# Windows 控制台常是 GBK，避免打印中文/符号时崩掉
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(errors="replace")
    except Exception:
        pass

KEY_NAME = ("姓名", "人名", "名字")
KEY_SID = ("学号", "学籍号", "考生号")
KEY_SCORE = ("评分", "分数", "得分")


def norm(v):
    if v is None:
        return ""
    return re.sub(r"\s+", "", str(v)).strip()


def pick(header, keys):
    for i, h in enumerate(header):
        h = norm(h)
        if h in keys:
            return i
    for i, h in enumerate(header):
        h = norm(h)
        if any(k in h for k in keys):
            return i
    return -1


def load_rows(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(c) if c is not None else "" for c in rows[0]]
    return header, rows[1:]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("output")
    ap.add_argument("--fail-max", type=float, default=2.0)
    ap.add_argument("--ref", default=None)
    ap.add_argument(
        "--exclude",
        default="",
        help="逗号分隔的姓名，整个跳过（例如缺考没来的人，既不算通过也不算不通过）",
    )
    a = ap.parse_args()

    header, rows = load_rows(a.input)
    i_name = pick(header, KEY_NAME)
    i_sid = pick(header, KEY_SID)
    i_score = pick(header, KEY_SCORE)
    if min(i_name, i_sid, i_score) < 0:
        sys.stderr.write(
            "认不出列。表头：%s\n需要能识别出 姓名 / 学号 / 评分 三类列。\n" % header
        )
        sys.exit(1)

    # 可选的学号回填表
    ref = {}
    if a.ref:
        r_header, r_rows = load_rows(a.ref)
        r_name = pick(r_header, KEY_NAME)
        r_sid = pick(r_header, KEY_SID)
        for r in r_rows:
            n, s = norm(r[r_name]), norm(r[r_sid])
            if n and re.match(r"^\d{6,20}$", s):
                ref[n] = s

    out, warned = [], []
    excluded = {norm(x) for x in a.exclude.replace("，", ",").split(",") if norm(x)}
    n_fail = n_pass = n_skip = 0
    for idx, r in enumerate(rows, start=2):
        name, sid = norm(r[i_name]), norm(r[i_sid])
        if not name and not sid:
            continue
        if name in excluded:
            n_skip += 1
            warned.append("第 %d 行 %s 按 --exclude 跳过（不计入通过/不通过）" % (idx, name))
            continue
        if not re.match(r"^\d{6,20}$", sid):
            if name in ref:
                warned.append("第 %d 行 %s 的学号在表里是 %r，已用参考表回填为 %s"
                              % (idx, name, str(r[i_sid]), ref[name]))
                sid = ref[name]
            else:
                warned.append("第 %d 行 %s 的学号 %r 不合法，且参考表里也没有，已按原样输出"
                              % (idx, name, str(r[i_sid])))

        raw = r[i_score]
        try:
            score = float(str(raw).strip())
        except (TypeError, ValueError):
            warned.append("第 %d 行 %s 的评分 %r 不是数字，按「未通过」处理" % (idx, name, raw))
            score = 0.0

        if score <= a.fail_max:
            res = "未通过"; n_fail += 1
        else:
            res = "通过"; n_pass += 1
        out.append((sid, name, res))

    with open(a.output, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["学号", "姓名", "结果"])
        w.writerows(out)

    print("写入 %s" % a.output)
    print("  总人数 %d   通过 %d   未通过 %d   排除 %d   （评分 <= %g 判为未通过）"
          % (len(out), n_pass, n_fail, n_skip, a.fail_max))
    for m in warned:
        print("  [注意] " + m)


if __name__ == "__main__":
    main()
